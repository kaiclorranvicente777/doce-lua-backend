import pandas as pd
import requests
import datetime
import os
import time
import win32com.client as win32
from PIL import ImageGrab
from openpyxl import load_workbook
import re

# === CONFIGURAÇÕES DO TELEGRAM ===
TOKEN = '8570909503:AAF8h_fee0KOSFg95TT6pOr_dDvTqlLQavk'
CHAT_ID = -1003380097906
URL = f'https://api.telegram.org/bot{TOKEN}/sendMessage'
GET_UPDATES_URL = f'https://api.telegram.org/bot{TOKEN}/getUpdates'

# === ARQUIVO DE LOG ===
log_path = r"C:\Users\Káic e Gabi\OneDrive\Documentos\log_execucao.txt"

def registrar_log(texto):
    with open(log_path, "a", encoding="utf-8") as log:
        log.write(f"{datetime.datetime.now().strftime('%d/%m/%Y %H:%M:%S')} - {texto}\n")

# === ARQUIVOS PRINCIPAIS ===
arquivo = r"C:\Users\Káic e Gabi\OneDrive\Documentos\Sistema.Central_Doce.Lua_11-2025.xlsx"
imagem_path = r"C:\Users\Káic e Gabi\OneDrive\Documentos\excel_intervalo.png"
aba_mensagens = "Mensagem-Excel"

# === GERA IMAGEM DO INTERVALO I1:AP27 ===
def gerar_imagem_excel():
    try:
        excel = win32.gencache.EnsureDispatch('Excel.Application')
        excel.Visible = True
        wb = excel.Workbooks.Open(arquivo)
        time.sleep(8)

        ws = wb.Sheets("Acompanhamento-Geral")
        ws.Activate()
        ws.Range("I1:AP27").Select()
        time.sleep(2)

        ws.Range("I1:AP27").CopyPicture(Appearance=1, Format=2)
        time.sleep(2)

        image = ImageGrab.grabclipboard()
        if image is None:
            raise Exception("Nenhuma imagem encontrada na área de transferência.")

        image.save(imagem_path)
        wb.Close(SaveChanges=False)
        excel.Quit()
        registrar_log("Imagem do calendário gerada com sucesso.")
    except Exception as e:
        registrar_log(f"Erro ao gerar imagem do calendário: {e}")

# === ENVIA IMAGEM DO CALENDÁRIO COM LEGENDA PERSONALIZADA ===
def enviar_imagem_calendario():
    try:
        with open(imagem_path, 'rb') as img:
            files = {'photo': img}
            data = {'chat_id': CHAT_ID, 'caption': 'Veja seu Calendário Programação Pedidos Doce Lua! 🤖🗓️'}
            response = requests.post(f'https://api.telegram.org/bot{TOKEN}/sendPhoto', files=files, data=data)
            registrar_log(f"Imagem do calendário enviada. Status: {response.status_code}")
        os.remove(imagem_path)
    except Exception as e:
        registrar_log(f"Erro ao enviar imagem do calendário: {e}")

# === LEITURA DA PLANILHA DE MENSAGENS ===
def ler_mensagens_planilha():
    try:
        df = pd.read_excel(arquivo, sheet_name=aba_mensagens, header=None)
        mensagens = df.iloc[0:30, 0].dropna().tolist()
        registrar_log(f"Planilha carregada com {len(mensagens)} mensagens.")
        return mensagens
    except Exception as e:
        registrar_log(f"Erro ao ler a planilha: {e}")
        return []

# === ENVIO DAS MENSAGENS COM CONTROLE DE FLOOD ===
def enviar_mensagens(mensagens):
    mensagens_enviadas = 0
    for i, mensagem in enumerate(mensagens, start=1):
        try:
            payload = {'chat_id': CHAT_ID, 'text': str(mensagem)}
            response = requests.post(URL, data=payload)
            if response.status_code == 200:
                registrar_log(f"[{i}] Mensagem enviada: {mensagem}")
                mensagens_enviadas += 1
            else:
                registrar_log(f"[{i}] Falha ao enviar: {mensagem} | Status: {response.status_code}")
        except Exception as e:
            registrar_log(f"[{i}] Erro ao enviar: {mensagem} | Erro: {e}")
        time.sleep(1)

    registrar_log(f"Total de mensagens enviadas com sucesso: {mensagens_enviadas} de {len(mensagens)}")

# === MENSAGEM FINAL ===
def enviar_mensagem_final():
    try:
        mensagem_final = "Que seu dia seja abençoado, adoçando as vidas das pessoas! 🤖💝"
        payload = {'chat_id': CHAT_ID, 'text': mensagem_final}
        response = requests.post(URL, data=payload)
        registrar_log(f"Mensagem final enviada: {mensagem_final} | Status: {response.status_code}")
    except Exception as e:
        registrar_log(f"Erro ao enviar mensagem final: {e}")

# === FECHA O EXCEL SE AINDA ESTIVER ABERTO ===
def fechar_excel_se_aberto():
    try:
        excel = win32.GetActiveObject("Excel.Application")
        excel.Quit()
        registrar_log("Excel fechado após envio das mensagens.")
    except Exception as e:
        registrar_log(f"Excel já estava fechado ou não foi possível encerrar: {e}")

# === VERIFICAÇÃO DE HORÁRIO PARA DESLIGAMENTO ===
def verificar_desligamento():
    hora_atual = datetime.datetime.now().time()
    hora_inicio = datetime.time(7, 0)
    hora_fim = datetime.time(8, 0)
    try:
        if hora_inicio <= hora_atual <= hora_fim:
            mensagem_desligar = "Estou deixando sua máquina ir descansar mais um pouco automaticamente! caso queira utilizar, ligue-a após as 8:00 🤖☕"
            requests.post(URL, data={'chat_id': CHAT_ID, 'text': mensagem_desligar})
            registrar_log("Mensagem de desligamento enviada.")
            os.system("shutdown /s /t 60")
        else:
            mensagem_ligada = "Sua máquina continuará ligada para que você trabalhe sem interrupções! Bom trabalho! 🤖🖥️"
            requests.post(URL, data={'chat_id': CHAT_ID, 'text': mensagem_ligada})
            registrar_log("Mensagem de permanência ligada enviada.")
    except Exception as e:
        registrar_log(f"Erro ao verificar hora ou enviar mensagem de desligamento: {e}")

# === UTILITÁRIOS TELEGRAM (POLLING) ===
def responder(chat_id, texto):
    try:
        requests.post(URL, data={'chat_id': chat_id, 'text': texto})
    except Exception as e:
        registrar_log(f"Erro ao responder no Telegram: {e}")

last_update_id = None
def ler_updates():
    global last_update_id
    try:
        params = {'timeout': 20}
        if last_update_id is not None:
            params['offset'] = last_update_id + 1
        resp = requests.get(GET_UPDATES_URL, params=params, timeout=30)
        data = resp.json()
        if not data.get('ok'):
            return []
        updates = data.get('result', [])
        if updates:
            last_update_id = updates[-1]['update_id']
        return updates
    except Exception as e:
        registrar_log(f"Erro em getUpdates: {e}")
        return []

def obter_texto(update):
    try:
        message = update.get('message') or update.get('channel_post') or update.get('edited_message')
        if not message:
            return None, None
        chat_id = message['chat']['id']
        texto = message.get('text')
        return chat_id, texto
    except Exception as e:
        registrar_log(f"Erro ao obter texto: {e}")
        return None, None

# === PARSER DO PEDIDO ===
def extrair_campos(texto):
    dados = {
        "Data": "", "Horario": "", "Nome": "", "Quantidade": "", "Produto": "",
        "UnidadeMedida": "", "SaborMassa": "", "PrepMassa": "",
        "SaborRecheio": "", "PrepRecheio": "", "DescricaoAdicional": "",
        "ValorAdicional": "", "ColetaAdicional": "", "Endereco": "", "Contato": "",
        "Entrega": "", "Pagamento": "", "RecebimentoFoto": "",
        "PublicacaoInstagram": "", "Preco": ""
    }

    item = re.search(r"🌙 Item:\s*(.+)", texto)
    if item:
        item_txt = item.group(1).strip()
        dados["Produto"] = item_txt
        qtd_un = re.search(r"(\d+)\s*(Kg|Kgs|Unidades?|Fatias?|Litros?|L|Cx|Caixas?|Pc?s?)", item_txt, re.IGNORECASE)
        if qtd_un:
            dados["Quantidade"] = qtd_un.group(1)
            dados["UnidadeMedida"] = qtd_un.group(2)

    # Sabor
    sabor = re.search(r"🌙 Sabor:\s*(.+)", texto)
    if sabor:
        dados["SaborRecheio"] = sabor.group(1).strip()
        # tentativa simples de extrair massa
        m_massa = re.search(r"massa\s+de\s+([^\)]+)", dados["SaborRecheio"], re.IGNORECASE)
        if m_massa:
            dados["SaborMassa"] = m_massa.group(1).strip()

    # Detalhes
    detalhes = re.search(r"🌙 Detalhes:\s*(.+)", texto)
    if detalhes:
        dados["DescricaoAdicional"] = detalhes.group(1).strip()

    # Adicional
    adicional = re.search(r"🌙 Adicional:\s*(.+)", texto)
    if adicional:
        dados["ValorAdicional"] = adicional.group(1).strip()

    # Endereço
    endereco = re.search(r"🌙 Endereço.*:\s*(.+)", texto)
    if endereco:
        dados["Endereco"] = endereco.group(1).strip()

    # Data
    data = re.search(r"🌙 Data:\s*(.+)", texto)
    if data:
        dados["Data"] = data.group(1).strip()

    # Total
    total = re.search(r"Total\s*\(R\$\s*([\d.,]+)\)", texto)
    if total:
        dados["Preco"] = total.group(1).strip()

    return dados

# === GRAVAÇÃO NA PLANILHA E MENSAGEM DE RETORNO ===
def gravar_excel(dados, arquivo, aba="Pedidos-via-bot"):
    wb = load_workbook(arquivo, data_only=True)  # <- ESSENCIAL
    ws = wb[aba]
    mensagem_retorno = ""

    # Encontrar primeira linha com AX == "#0000"
    for row in range(2, ws.max_row + 1):
        valor_ax = str(ws[f"AX{row}"].value).strip()
        if valor_ax == "#0000":
            ws[f"E{row}"] = dados["Data"]
            ws[f"F{row}"] = dados["Horario"]
            ws[f"G{row}"] = dados["Nome"]
            ws[f"H{row}"] = dados["Quantidade"]
            ws[f"I{row}"] = dados["UnidadeMedida"]
            ws[f"J{row}"] = dados["Produto"]
            ws[f"L{row}"] = dados["SaborMassa"]
            ws[f"M{row}"] = dados["PrepMassa"]
            ws[f"N{row}"] = dados["SaborRecheio"]
            ws[f"O{row}"] = dados["PrepRecheio"]
            ws[f"P{row}"] = dados["DescricaoAdicional"]
            ws[f"Q{row}"] = dados["ValorAdicional"]
            ws[f"R{row}"] = dados["ColetaAdicional"]
            ws[f"S{row}"] = dados["Endereco"]
            ws[f"T{row}"] = dados["Contato"]
            ws[f"U{row}"] = dados["Entrega"]
            ws[f"V{row}"] = dados["Pagamento"]
            ws[f"W{row}"] = dados["RecebimentoFoto"]
            ws[f"X{row}"] = dados["PublicacaoInstagram"]
            ws[f"Y{row}"] = dados["Preco"]

            codigo_pedido = valor_ax
            status_atual = ws[f"D{row}"].value
            mensagem_retorno = f"Pedido {codigo_pedido} Adicionado ao sistema! ✅"
            if status_atual:
                mensagem_retorno += f" | Status atual do pedido: {status_atual}"
            break

    wb.save(arquivo)
    return mensagem_retorno

# === LOOP PRINCIPAL DE ESCUTA DE MENSAGENS ===
def modo_escuta_pedidos():
    registrar_log("Iniciando modo de escuta de pedidos...")
    esperando_pedido = False

    while True:
        updates = ler_updates()
        for u in updates:
            chat_id, texto = obter_texto(u)
            if chat_id != CHAT_ID or not texto:
                continue

            if "adicionar pedido ao sistema" in texto.strip().lower():
                esperando_pedido = True
                responder(CHAT_ID, "Ok, envie agora a mensagem de confirmação do pedido 🌙 ...")
                registrar_log("Flag esperando_pedido ativada.")
                continue

            if esperando_pedido and texto.startswith("Seu pedido Doce Lua confirmado:"):
                fechar_excel_se_aberto()
                dados = extrair_campos(texto)
                retorno = gravar_excel(dados, arquivo, aba="Pedidos-via-bot")
                responder(CHAT_ID, retorno if retorno else "Pedido registrado. ✅")
                registrar_log(f"Pedido gravado. Retorno: {retorno}")
                esperando_pedido = False
                continue

        time.sleep(1)

# === EXECUÇÃO SEQUENCIAL ===
if __name__ == "__main__":
    gerar_imagem_excel()
    enviar_imagem_calendario()
    mensagens = ler_mensagens_planilha()
    enviar_mensagens(mensagens)
    enviar_mensagem_final()
    fechar_excel_se_aberto()
    verificar_desligamento()
    modo_escuta_pedidos()
