from openpyxl import load_workbook

# Caminho do arquivo Excel
arquivo = "Sistema.Central_Doce.Lua_11-2025.xlsx"   # ajuste se a extensão for diferente
aba = "Pedidos-via-bot"

wb = load_workbook(arquivo, data_only=False)
ws = wb[aba]

print("=== Teste de leitura da coluna AX ===")
for row in range(2, ws.max_row + 1):
    valor_bruto = ws[f"AX{row}"].value
    valor_convertido = str(valor_bruto).strip() if valor_bruto is not None else "None"
    print(f"Linha {row} -> AX bruto: {valor_bruto} | AX convertido: {valor_convertido}")
